"""Plantilla de Excel para cargar transacciones en lote (descarga + lectura).

Filosofía: este módulo sólo traduce filas de una hoja de cálculo a los
mismos campos que ya espera `TransactionSerializer` (o produce un `RowError`
si no puede) -- las reglas de negocio de verdad (cartera privada ajena,
categoría del tipo que no corresponde, workspace cruzado, etc.) las sigue
validando ese serializer, que es a quien `TransactionViewSet.import_xlsx`
le pasa el resultado de `parse_workbook`. Así hay un solo lugar que decide
qué es una transacción válida.
"""
from __future__ import annotations

import datetime as dt
import io
import zipfile
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from apps.transactions.models import Transaction

HEADERS = [
    "fecha",
    "tipo",
    "categoria",
    "cartera",
    "cartera_destino",
    "monto",
    "nota",
    "cuenta_presupuesto",
]

# Mismas etiquetas en español que ya usa la exportación a CSV
# (`src/lib/export.ts` del frontend) -- para poder exportar, editar en Excel
# y volver a importar sin traducir nada a mano.
TYPE_LABELS = {
    Transaction.TYPE_INCOME: "ingreso",
    Transaction.TYPE_EXPENSE: "gasto",
    Transaction.TYPE_TRANSFER: "transferencia",
}
TYPE_BY_LABEL = {label: code for code, label in TYPE_LABELS.items()}

EXAMPLE_ROWS = [
    ["2026-09-01", "gasto", "Supermercado", "Cuenta principal", "", "45.50", "Compras de la semana", "si"],
    ["2026-09-02", "ingreso", "Sueldo", "Cuenta principal", "", "1200.00", "", "si"],
    ["2026-09-03", "transferencia", "", "Cuenta principal", "Ahorro", "200.00", "Aporte a ahorro", ""],
]

INSTRUCTIONS = [
    "Cómo llenar la hoja «Transacciones»",
    "",
    "fecha: formato AAAA-MM-DD (p. ej. 2026-09-01).",
    "tipo: gasto, ingreso o transferencia.",
    "categoria: nombre EXACTO de una categoría de este presupuesto (ver hoja «Categorías»). Vacío en transferencias.",
    "cartera: nombre EXACTO de la cartera de origen (ver hoja «Carteras»).",
    "cartera_destino: solo en transferencias -- nombre EXACTO de la cartera que recibe.",
    "monto: número mayor que 0 (p. ej. 45.50).",
    "nota: texto libre, opcional.",
    "cuenta_presupuesto: «si» o «no» -- si el gasto cuenta contra el presupuesto de su categoría. Opcional, «si» por defecto.",
    "",
    "Borra las 3 filas de ejemplo (en gris) antes de importar -- si las dejás, se importan como transacciones reales.",
    "Los nombres de categoría y cartera tienen que coincidir tal cual (sin importar mayúsculas) con los de sus hojas.",
]


class InvalidWorkbook(Exception):
    """El archivo subido no es un .xlsx legible."""


class RowError(Exception):
    """Una fila de la plantilla no se puede traducir a una transacción."""


def _style_header(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F6F4F")
    cell.alignment = Alignment(horizontal="center")


def build_template(wallets, categories) -> bytes:
    """Genera el .xlsx de la plantilla, con los nombres reales de carteras y
    categorías del workspace en hojas de referencia (para que el usuario
    sepa qué escribir sin adivinar)."""
    wb = Workbook()

    sheet = wb.active
    sheet.title = "Transacciones"
    for col, name in enumerate(HEADERS, start=1):
        _style_header(sheet.cell(row=1, column=col, value=name))

    example_font = Font(italic=True, color="808080")
    for r, row in enumerate(EXAMPLE_ROWS, start=2):
        for c, value in enumerate(row, start=1):
            sheet.cell(row=r, column=c, value=value).font = example_font

    for col, width in enumerate([12, 14, 22, 22, 22, 12, 30, 16], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A2"

    notes = wb.create_sheet("Instrucciones")
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = notes.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 100

    wallets_sheet = wb.create_sheet("Carteras")
    _style_header(wallets_sheet.cell(row=1, column=1, value="cartera"))
    _style_header(wallets_sheet.cell(row=1, column=2, value="moneda"))
    for r, wallet in enumerate(wallets, start=2):
        wallets_sheet.cell(row=r, column=1, value=wallet.name)
        wallets_sheet.cell(row=r, column=2, value=wallet.currency)
    wallets_sheet.column_dimensions["A"].width = 28
    wallets_sheet.column_dimensions["B"].width = 10

    categories_sheet = wb.create_sheet("Categorías")
    _style_header(categories_sheet.cell(row=1, column=1, value="categoria"))
    _style_header(categories_sheet.cell(row=1, column=2, value="tipo"))
    _style_header(categories_sheet.cell(row=1, column=3, value="grupo"))
    for r, category in enumerate(categories, start=2):
        categories_sheet.cell(row=r, column=1, value=category.name)
        categories_sheet.cell(row=r, column=2, value=TYPE_LABELS.get(category.type, category.type))
        categories_sheet.cell(row=r, column=3, value=category.parent.name if category.parent_id else "")
    categories_sheet.column_dimensions["A"].width = 26
    categories_sheet.column_dimensions["B"].width = 12
    categories_sheet.column_dimensions["C"].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_is_blank(values: tuple) -> bool:
    return all(_cell_text(v) == "" for v in values)


def _parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _cell_text(value)
    try:
        return dt.date.fromisoformat(text)
    except ValueError:
        raise RowError(f"fecha inválida: «{value}» (usa AAAA-MM-DD).")


def _parse_amount(value):
    text = _cell_text(value)
    if not text:
        raise RowError("monto requerido.")
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        amount = Decimal(str(value))
    else:
        # Tolera coma decimal (Excel en español) si no hay punto ya.
        normalized = text.replace(",", ".") if "." not in text else text
        try:
            amount = Decimal(normalized)
        except InvalidOperation:
            raise RowError(f"monto inválido: «{value}».")
    if amount <= 0:
        raise RowError("monto tiene que ser mayor que 0.")
    return amount


def _parse_counts_toward_budget(value) -> bool:
    text = _cell_text(value).lower()
    return text not in ("no", "false", "0")


def _index_by_name(items, key) -> dict:
    """`key(item)` (ya normalizada por el caller) -> objeto, o `None` si hay
    más de uno con esa clave (ambigüedad que se reporta fila por fila, en
    vez de adivinar cuál de los dos era)."""
    index: dict = {}
    for item in items:
        k = key(item)
        index[k] = None if k in index else item
    return index


def index_wallets_by_name(wallets) -> dict:
    return _index_by_name(wallets, lambda w: w.name.strip().lower())


def index_categories_by_name(categories) -> dict:
    """Clave `(nombre en minúsculas, tipo)`: el mismo nombre puede existir
    como categoría de ingreso y de gasto a la vez sin ser ambiguo, porque el
    tipo de la fila (columna `tipo`) ya desambigua cuál de las dos es."""
    return _index_by_name(categories, lambda c: (c.name.strip().lower(), c.type))


def row_to_payload(values: tuple, wallets_by_name: dict, categories_by_name: dict) -> dict:
    """`values` son las 8 celdas de una fila (orden de `HEADERS`). Devuelve
    el dict listo para `TransactionSerializer(data=...)`, o lanza `RowError`."""
    fecha, tipo, categoria, cartera, cartera_destino, monto, nota, cuenta_presupuesto = values

    type_label = _cell_text(tipo).lower()
    txn_type = TYPE_BY_LABEL.get(type_label)
    if txn_type is None:
        raise RowError(f"tipo inválido: «{tipo}» (usa gasto, ingreso o transferencia).")

    date = _parse_date(fecha)
    amount = _parse_amount(monto)

    wallet_name = _cell_text(cartera)
    if not wallet_name:
        raise RowError("cartera requerida.")
    wallet = wallets_by_name.get(wallet_name.lower())
    if wallet is None:
        raise RowError(f"no existe una cartera llamada «{wallet_name}» (o hay más de una con ese nombre).")

    payload = {
        "type": txn_type,
        "wallet": str(wallet.id),
        "amount": str(amount),
        "date": date.isoformat(),
        "description": _cell_text(nota),
        "counts_toward_budget": _parse_counts_toward_budget(cuenta_presupuesto),
        "source": Transaction.SOURCE_EXCEL_IMPORT,
    }

    if txn_type == Transaction.TYPE_TRANSFER:
        to_name = _cell_text(cartera_destino)
        if not to_name:
            raise RowError("cartera_destino requerida en una transferencia.")
        to_wallet = wallets_by_name.get(to_name.lower())
        if to_wallet is None:
            raise RowError(
                f"no existe una cartera llamada «{to_name}» (o hay más de una con ese nombre)."
            )
        payload["to_wallet"] = str(to_wallet.id)
    else:
        category_name = _cell_text(categoria)
        if not category_name:
            raise RowError("categoria requerida en ingresos y gastos.")
        category = categories_by_name.get((category_name.lower(), txn_type))
        if category is None:
            raise RowError(
                f"no existe una categoría de {TYPE_LABELS[txn_type]} llamada «{category_name}» "
                "(o hay más de una con ese nombre)."
            )
        payload["category"] = str(category.id)

    return payload


def parse_workbook(file) -> list[tuple[int, tuple]]:
    """Lee el .xlsx subido y devuelve `(número de fila, celdas)` por cada
    fila con datos de la primera hoja, saltando el encabezado y las filas
    vacías. No valida contenido -- eso es cosa de `row_to_payload`."""
    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise InvalidWorkbook(str(exc))

    sheet = wb.worksheets[0]
    rows = []
    for i, values in enumerate(sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True), start=2):
        padded = tuple(values) + (None,) * (len(HEADERS) - len(values))
        if _row_is_blank(padded):
            continue
        rows.append((i, padded))
    return rows
