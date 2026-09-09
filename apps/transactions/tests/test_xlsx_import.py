"""Importar transacciones en lote desde la plantilla de Excel (.xlsx), y
descargar esa plantilla ya con las carteras/categorías del workspace."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.accounts.models import Wallet
from apps.transactions.models import Category, Transaction
from apps.workspaces.models import Membership, Workspace

User = get_user_model()
HEADER = "HTTP_X_WORKSPACE_ID"

HEADERS = ["fecha", "tipo", "categoria", "cartera", "cartera_destino", "monto", "nota", "cuenta_presupuesto"]


def _xlsx_upload(rows, name="movimientos.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(
        name,
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class XlsxImportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("alice", "a@example.com", "pw")
        cls.outsider = User.objects.create_user("mallory", "m@example.com", "pw")
        cls.ws = Workspace.objects.create(name="Casa")
        Membership.objects.create(workspace=cls.ws, user=cls.user, role=Membership.ROLE_OWNER)

        cls.checking = Wallet.objects.create(
            workspace=cls.ws, name="Cuenta principal", purpose=Wallet.PURPOSE_SPENDING,
            opening_balance=Decimal("500.00"),
        )
        cls.savings = Wallet.objects.create(
            workspace=cls.ws, name="Ahorro", purpose=Wallet.PURPOSE_SAVINGS,
        )
        cls.food = Category.objects.create(
            workspace=cls.ws, name="Supermercado", type=Category.TYPE_EXPENSE
        )
        cls.salary = Category.objects.create(
            workspace=cls.ws, name="Sueldo", type=Category.TYPE_INCOME
        )

    def setUp(self):
        self.client.force_authenticate(self.user)

    def _import(self, rows, filename="movimientos.xlsx"):
        return self.client.post(
            "/api/v1/transactions/import/",
            {"file": _xlsx_upload(rows, filename)},
            format="multipart",
            **{HEADER: str(self.ws.id)},
        )

    # -- descarga de la plantilla -------------------------------------
    def test_template_download_includes_real_wallets_and_categories(self):
        res = self.client.get(
            "/api/v1/transactions/import-template/", **{HEADER: str(self.ws.id)}
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        content = b"".join(res.streaming_content) if res.streaming else res.content
        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb["Transacciones"]["A1"].value, "fecha")

        wallet_names = {row[0].value for row in wb["Carteras"].iter_rows(min_row=2)}
        self.assertEqual(wallet_names, {"Cuenta principal", "Ahorro"})

        category_names = {row[0].value for row in wb["Categorías"].iter_rows(min_row=2)}
        self.assertEqual(category_names, {"Supermercado", "Sueldo"})

    # -- import feliz ---------------------------------------------------
    def test_imports_expense_income_and_transfer(self):
        res = self._import(
            [
                ["2026-09-01", "gasto", "Supermercado", "Cuenta principal", "", "45.50", "Compras", "si"],
                ["2026-09-02", "ingreso", "Sueldo", "Cuenta principal", "", "1200.00", "", "si"],
                ["2026-09-03", "transferencia", "", "Cuenta principal", "Ahorro", "200.00", "Aporte", ""],
            ]
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertEqual(res.data["created"], 3)
        self.assertEqual(res.data["errors"], [])
        self.assertEqual(Transaction.objects.count(), 3)

        imported = Transaction.objects.filter(source=Transaction.SOURCE_EXCEL_IMPORT)
        self.assertEqual(imported.count(), 3)

        self.checking.refresh_from_db()
        self.savings.refresh_from_db()
        # 500 - 45.50 + 1200 - 200 = 1454.50
        self.assertEqual(self.checking.current_balance, Decimal("1454.50"))
        self.assertEqual(self.savings.current_balance, Decimal("200.00"))

    def test_ignores_blank_rows(self):
        res = self._import(
            [
                ["2026-09-01", "gasto", "Supermercado", "Cuenta principal", "", "10.00", "", "si"],
                ["", "", "", "", "", "", "", ""],
            ]
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["created"], 1)
        self.assertEqual(res.data["errors"], [])

    # -- errores fila por fila, sin frenar al resto ----------------------
    def test_partial_success_reports_errors_per_row_without_blocking_valid_ones(self):
        res = self._import(
            [
                ["2026-09-01", "gasto", "Supermercado", "Cuenta principal", "", "10.00", "", "si"],
                ["fecha-mala", "gasto", "Supermercado", "Cuenta principal", "", "10.00", "", "si"],
                ["2026-09-02", "gasto", "Categoria que no existe", "Cuenta principal", "", "10.00", "", "si"],
                ["2026-09-03", "gasto", "Supermercado", "Cartera que no existe", "", "10.00", "", "si"],
                ["2026-09-04", "gasto", "Supermercado", "Cuenta principal", "", "0", "", "si"],
                ["2026-09-05", "algo-raro", "Supermercado", "Cuenta principal", "", "10.00", "", "si"],
            ]
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertEqual(res.data["created"], 1)
        self.assertEqual(len(res.data["errors"]), 5)
        # Los números de fila son los de la hoja real (2 = primera fila de datos).
        self.assertEqual([e["row"] for e in res.data["errors"]], [3, 4, 5, 6, 7])

    def test_transfer_requires_to_wallet(self):
        res = self._import(
            [["2026-09-01", "transferencia", "", "Cuenta principal", "", "10.00", "", ""]]
        )
        self.assertEqual(res.data["created"], 0)
        self.assertEqual(len(res.data["errors"]), 1)
        self.assertIn("cartera_destino", res.data["errors"][0]["message"])

    def test_duplicate_wallet_name_is_reported_as_ambiguous_not_guessed(self):
        Wallet.objects.create(workspace=self.ws, name="Cuenta principal", purpose=Wallet.PURPOSE_SPENDING)
        res = self._import(
            [["2026-09-01", "gasto", "Supermercado", "Cuenta principal", "", "10.00", "", "si"]]
        )
        self.assertEqual(res.data["created"], 0)
        self.assertEqual(len(res.data["errors"]), 1)

    def test_rejects_non_xlsx_file(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        res = self.client.post(
            "/api/v1/transactions/import/",
            {"file": SimpleUploadedFile("notas.txt", b"esto no es un excel", content_type="text/plain")},
            format="multipart",
            **{HEADER: str(self.ws.id)},
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_rejects_file_too_large(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        big = SimpleUploadedFile(
            "movimientos.xlsx", b"0" * (6 * 1024 * 1024),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        res = self.client.post(
            "/api/v1/transactions/import/",
            {"file": big},
            format="multipart",
            **{HEADER: str(self.ws.id)},
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_requires_file(self):
        res = self.client.post(
            "/api/v1/transactions/import/", {}, format="multipart", **{HEADER: str(self.ws.id)}
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    # -- privacidad: no se puede importar contra una cartera privada ajena
    def test_cannot_import_against_another_users_private_wallet(self):
        Membership.objects.create(workspace=self.ws, user=self.outsider, role=Membership.ROLE_MEMBER)
        private = Wallet.objects.create(
            workspace=self.ws, name="Secreta", purpose=Wallet.PURPOSE_SPENDING,
            visibility=Wallet.VISIBILITY_PRIVATE, owner=self.outsider,
        )
        res = self._import(
            [["2026-09-01", "gasto", "Supermercado", private.name, "", "10.00", "", "si"]]
        )
        self.assertEqual(res.data["created"], 0)
        self.assertEqual(len(res.data["errors"]), 1)
        self.assertFalse(Transaction.objects.filter(wallet=private).exists())
